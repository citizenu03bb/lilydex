import pandas as pd
import pytest
from pokeprice.report import (
    detect_card_anomaly,
    apply_pricing_and_anomalies,
    CONDITION_MULTIPLIERS,
)


def test_condition_multipliers():
    assert CONDITION_MULTIPLIERS["NM"] == 1.0
    assert CONDITION_MULTIPLIERS["EX"] == 0.85
    assert CONDITION_MULTIPLIERS["GD"] == 0.70
    assert CONDITION_MULTIPLIERS["LP"] == 0.55
    assert CONDITION_MULTIPLIERS["PL"] == 0.40
    assert CONDITION_MULTIPLIERS["PO"] == 0.15


def test_detect_counterfeit_risk_base_set_non_holo():
    # Base set 6/102 (Léviator) non-holo should trigger anomaly
    row = {
        "Identified_Name": "Léviator",
        "Extension": "Set de Base",
        "Numero_Officiel": "6/102",
        "Is_Foil": "Non",
    }
    warning = detect_card_anomaly(row)
    assert warning != ""
    assert "Incohérence Holo" in warning

    # Base set 6/102 holo should NOT trigger anomaly
    row_holo = {
        "Identified_Name": "Léviator",
        "Extension": "Set de Base",
        "Numero_Officiel": "6/102",
        "Is_Foil": "Oui",
    }
    assert detect_card_anomaly(row_holo) == ""


def test_apply_pricing_and_anomalies():
    data = [
        {
            "Page": 1,
            "Position": 1,
            "Identified_Name": "Pikachu",
            "Extension": "EX Fantômes Holon",
            "Numero_Officiel": "78/110",
            "Is_Foil": "Oui",
            "Etat": "PO",
            "Prix_Reference_EUR": 100.0,
        },
        {
            "Page": 1,
            "Position": 2,
            "Identified_Name": "Léviator",
            "Extension": "Set de Base",
            "Numero_Officiel": "6/102",
            "Is_Foil": "Non",
            "Etat": "LP",
            "Prix_Reference_EUR": 20.0,
        },
        {
            "Page": 1,
            "Position": 3,
            "Identified_Name": "Dracaufeu",
            "Extension": "Set de Base",
            "Numero_Officiel": "4/102",
            "Is_Foil": "Oui",
            "Etat": "NM",
            "Prix_Reference_EUR": 250.0,
        },
    ]
    df = pd.DataFrame(data)
    df_res = apply_pricing_and_anomalies(df)

    # Pikachu PO -> 100 * 0.15 = 15.00
    assert df_res.loc[0, "Prix_Simule_EUR"] == 15.0
    assert df_res.loc[0, "Est_Suspecte"] == "Non"

    # Holo uncertainty needs review; do not write off the card as counterfeit.
    assert df_res.loc[1, "Est_Suspecte"] == "Oui"
    assert pd.isna(df_res.loc[1, "Prix_Simule_EUR"])

    # Dracaufeu NM -> 250 * 1.0 = 250.0
    assert df_res.loc[2, "Prix_Simule_EUR"] == 250.0
    assert df_res.loc[2, "Est_Suspecte"] == "Non"


def test_wailord_200hp_not_counterfeit():
    row = {
        "Identified_Name": "Wailord ex",
        "Extension": "EX Tempête de Sable",
        "Numero_Officiel": "100/100",
        "PV": 200,
        "PV_Officiel": 200,
        "Is_Foil": "Oui",
        "Etat": "NM",
        "Prix_Reference_EUR": 100.0,
    }
    df = apply_pricing_and_anomalies(pd.DataFrame([row]))
    assert df.loc[0, "Est_Suspecte"] == "Non"
    assert df.loc[0, "Prix_Simule_EUR"] == 100


@pytest.mark.parametrize("grade,price", [("", None), ("NM/EX", 8.5), ("PO", 1.5)])
def test_grade_discount_and_unknown(grade, price):
    df = apply_pricing_and_anomalies(
        pd.DataFrame([{"Etat": grade, "Prix_Reference_EUR": 10}])
    )
    assert (
        pd.isna(df.loc[0, "Prix_Simule_EUR"])
        if price is None
        else df.loc[0, "Prix_Simule_EUR"] == price
    )


def test_unresolved_and_missing_prices_excluded():
    df = apply_pricing_and_anomalies(
        pd.DataFrame(
            [
                {
                    "Etat": "NM",
                    "Prix_Reference_EUR": 10,
                    "Statut_Identification": "review",
                },
                {
                    "Etat": "NM",
                    "Prix_Reference_EUR": None,
                    "Statut_Identification": "confirmed",
                },
            ]
        )
    )
    assert df.Prix_Simule_EUR.isna().all()
    assert pd.isna(df.loc[1, "Prix_Reference_EUR"])


def test_safe_exports_and_layout(tmp_path):
    from pokeprice.report import export_excel, export_csv, read_report
    from openpyxl import load_workbook

    df = apply_pricing_and_anomalies(
        pd.DataFrame(
            [
                {
                    "OCR_Name": "=1+1",
                    "Identified_Name": "Pikachu",
                    "Etat": "EX",
                    "Defauts": "long description " * 40,
                    "Prix_Reference_EUR": 10.0,
                    "Page": 1,
                    "Position": 1,
                }
            ]
        )
    )
    excel = tmp_path / "safe.xlsx"
    csv = tmp_path / "safe.csv"
    export_excel(df, excel)
    export_csv(df, csv)
    ws = load_workbook(excel).active
    assert not any(c.data_type == "f" for row in ws for c in row)
    assert ws.freeze_panes == "D2" and ws.auto_filter.ref
    assert max(d.width for d in ws.column_dimensions.values()) <= 36
    assert read_report(csv).loc[0, "OCR_Name"] == "=1+1"
    assert read_report(excel).loc[0, "OCR_Name"] == "=1+1"
    price_cell = next(c for c in ws[1] if c.value == "Prix_Simule_EUR")
    assert ws.cell(2, price_cell.column).value == 8.5
    assert "€" in ws.cell(2, price_cell.column).number_format
    assert ws.column_dimensions[price_cell.column_letter].hidden
    assert "Etat_A_Confirmer" in [c.value for c in ws[1]]


def test_legacy_reports_migrate_and_simulation_is_not_compounded(tmp_path):
    from pokeprice.report import read_report, export_csv

    original = pd.DataFrame(
        [
            {
                "Etat": "PO",
                "Prix_Estime_EUR": 100.0,
                "Prix_Catalogue_EUR": 100.0,
                "Prix_Ajuste_EUR": 15.0,
                "Decote_Etat": "-85% (PO)",
            }
        ]
    )
    path = tmp_path / "old.csv"
    original.to_csv(path, index=False)
    df = apply_pricing_and_anomalies(read_report(path))
    assert (
        df.loc[0, "Prix_Reference_EUR"] == 100.0
        and df.loc[0, "Prix_Simule_EUR"] == 15.0
    )
    assert df.loc[0, "Coefficient_Simulation_Etat"] == 0.15
    assert df.loc[0, "Etat_A_Confirmer"] == "Oui"
    assert df.loc[0, "Source_Etat"] == "Inconnue (historique)"
    assert df.loc[0, "Statut_Revision"] == "Simulation disponible"
    assert not any(
        c in df
        for c in (
            "Prix_Estime_EUR",
            "Prix_Catalogue_EUR",
            "Prix_Ajuste_EUR",
            "Decote_Etat",
        )
    )
    export_csv(df, path)
    reloaded = apply_pricing_and_anomalies(read_report(path))
    pd.testing.assert_frame_equal(df, reloaded, check_dtype=False)


def test_terminal_summary_shows_references_instead_of_sale_estimate(capsys):
    from pokeprice.report import print_summary

    df = apply_pricing_and_anomalies(
        pd.DataFrame(
            [
                {
                    "Page": 1,
                    "Position": 1,
                    "Etat": "PO",
                    "Is_Foil": "Non",
                    "Identified_Name": "Example",
                    "Prix_Reference_EUR": 100.0,
                    "Statut_Identification": "confirmed",
                }
            ]
        )
    )
    print_summary(df)
    text = capsys.readouterr().out
    assert "100.00 €" in text and "SOMME DES RÉFÉRENCES" in text
    assert "ESTIMATION" not in text and "décote globale" not in text
